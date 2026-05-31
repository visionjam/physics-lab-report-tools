#!/usr/bin/env python3
"""
Reference script for docx-format-standardizer skill.
Processes a target .docx to match a reference document's formatting.

Key patterns demonstrated:
- Remove non-standard heading numbering (r'^\d+(?:\.\d+)+\s+')
- Set run-level fonts (with east-asia support for Chinese)
- Handle 装订线, 封面信息, 思考题 zones
"""

import re
from docx import Document
from docx.shared import Emu
from docx.oxml.ns import qn

# EMU values (English Metric Units)
PT12 = 152400   # 正文 12pt
PT15 = 190500   # 二级标题 / 思考题 15pt
PT18 = 228600   # 一级标题 18pt
PT36 = 457200   # 封面标题 36pt

FONT_HEITI = "黑体"
FONT_TNR = "Times New Roman"

HEADING_NUM_PATTERN = re.compile(r'^\d+(?:\.\d+)+\s+')
QUESTION_PATTERN = re.compile(r'^\d+\.\s+')


def set_run_font(run, font_name=None, font_size=None, bold=None):
    """Set font properties on a run, including Chinese font support."""
    if font_name is not None:
        run.font.name = font_name
        rpr = run._element.get_or_add_rPr()
        rFonts = rpr.find(qn('w:rFonts'))
        if rFonts is None:
            rFonts = run._element.makeelement(qn('w:rFonts'), {})
            rpr.insert(0, rFonts)
        rFonts.set(qn('w:eastAsia'), font_name)
    if font_size is not None:
        run.font.size = Emu(font_size)
    if bold is not None:
        run.font.bold = bold


def set_all_runs_font(paragraph, font_name=None, font_size=None, bold=None):
    for run in paragraph.runs:
        set_run_font(run, font_name, font_size, bold)


def remove_numbering(text):
    return HEADING_NUM_PATTERN.sub('', text)


def main():
    import sys
    if len(sys.argv) < 2:
        print("Usage: python script.py target.docx [output.docx]")
        return

    src = sys.argv[1]
    out = sys.argv[2] if len(sys.argv) > 2 else src.replace('.docx', '_格式化版.docx')

    doc = Document(src)

    thinking_section = False
    question_mode = False

    for i, p in enumerate(doc.paragraphs):
        text = p.text.strip()
        if not text:
            continue

        style_name = p.style.name

        # Detect 思考题 section
        if text == '思考题':
            thinking_section = True
            question_mode = False
            continue

        # Detect end of 思考题 section (装订线)
        if thinking_section and text in ('装', '订') and \
           p.alignment is not None and 'CENTER' in str(p.alignment):
            thinking_section = False
            question_mode = False
            continue

        # Handle 思考题 questions
        if thinking_section:
            if QUESTION_PATTERN.match(text):
                set_all_runs_font(p, FONT_HEITI, PT15, True)
                question_mode = True
                continue
            elif text.startswith('答'):
                set_all_runs_font(p, None, PT12, None)
                question_mode = False
                continue
            if question_mode:
                set_all_runs_font(p, None, PT12, None)
                continue

        # Handle section headings (一级标题)
        HEADINGS = ['实验目的', '实验原理', '实验内容', '实验仪器',
                     '数据记录处理', '讨论与分析', '实验结论']
        if style_name == 'Normal' and text in HEADINGS:
            set_all_runs_font(p, FONT_HEITI, PT18, True)
            continue

        # Handle 实验名称 page heading (centered, 36pt)
        if text == '实验名称' and p.alignment is not None and 'CENTER' in str(p.alignment):
            set_all_runs_font(p, FONT_HEITI, PT36, True)
            continue

        # Handle Heading 3 / 二级标题 — remove numbering, set 15pt TNR bold
        if style_name == 'Heading 3':
            if HEADING_NUM_PATTERN.match(p.text):
                new_text = remove_numbering(p.text)
                for run in p.runs:
                    run.text = ''
                if p.runs:
                    p.runs[0].text = new_text
                else:
                    p.add_run(new_text)
            set_all_runs_font(p, FONT_TNR, PT15, True)
            continue

        # Handle Heading 4 / 三级标题
        if style_name == 'Heading 4':
            if HEADING_NUM_PATTERN.match(p.text):
                new_text = remove_numbering(p.text)
                for run in p.runs:
                    run.text = ''
                if p.runs:
                    p.runs[0].text = new_text
                else:
                    p.add_run(new_text)
            set_all_runs_font(p, FONT_TNR, PT15, True)
            continue

        # Body text (Normal) — set 12pt if not already
        if style_name == 'Normal':
            # Skip cover page info (already correct)
            has_18pt = any(r.font.size and r.font.size == PT18 for r in p.runs if r.font.size)
            has_36pt = any(r.font.size and r.font.size == PT36 for r in p.runs if r.font.size)
            if has_18pt or has_36pt:
                continue
            has_size = any(r.font.size for r in p.runs if r.font.size)
            if not has_size:
                set_all_runs_font(p, None, PT12, None)

    # Process tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    has_size = any(r.font.size for r in p.runs if r.font.size)
                    if not has_size:
                        set_all_runs_font(p, None, PT12, None)

    doc.save(out)
    print(f"Done: {out}")


if __name__ == '__main__':
    main()
